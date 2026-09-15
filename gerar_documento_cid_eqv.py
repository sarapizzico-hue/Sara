#!/usr/bin/env python3
"""Gera a proposta CID (Word HTML + .doc + .docx) no formato G6/Google Docs."""

from html import escape
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

RED = "A10F14"
RED2 = "E50914"
INK = "1A1A1A"
MUT = "4A4A4A"
CREAM = "F7F5F4"
LINE = "E6D4D4"
WHITE = "FFFFFF"
SOFT = "F8EFEF"
DARK = "280001"

OUT_DIR = Path("/workspace/cid")
OUT_HTML = OUT_DIR / "documento.html"
OUT_DOC = OUT_DIR / "CID-Empresa-que-Vende.doc"
OUT_DOCX = OUT_DIR / "CID-Empresa-que-Vende.docx"


def set_run_font(run, name="Calibri", size=11, bold=False, color=INK, italic=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def shade_cell(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_borders(cell, color="E6D4D4", sz="4"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), sz)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        tcBorders.append(el)
    tcPr.append(tcBorders)


def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for m, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        node = OxmlElement(f"w:{m}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def set_table_width(table, width_cm):
    table.autofit = False
    table.allow_autofit = False
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
    tblW = OxmlElement("w:tblW")
    tblW.set(qn("w:w"), str(int(width_cm * 567)))
    tblW.set(qn("w:type"), "dxa")
    tblPr.append(tblW)


def prevent_row_split(row):
    tr = row._tr
    trPr = tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    trPr.append(cant)


def cell_para(cell, text, size=11, bold=False, color=INK, align="left", space_after=4):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }[align]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    return p


def add_cell_run(cell, text, size=11, bold=False, color=INK, italic=False):
    p = (
        cell.paragraphs[0]
        if cell.paragraphs[0].text == "" and len(cell.paragraphs) == 1
        else cell.add_paragraph()
    )
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color, italic=italic)
    return p


def add_p(doc, text, size=11, bold=False, color=INK, space_before=0, space_after=8, align="left", italic=False):
    p = doc.add_paragraph()
    p.alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }[align]
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color, italic=italic)
    return p


def add_heading_styled(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16 if level == 1 else 12)
    p.paragraph_format.space_after = Pt(8)
    if level == 1:
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "12")
        bottom.set(qn("w:space"), "4")
        bottom.set(qn("w:color"), RED)
        pBdr.append(bottom)
        pPr.append(pBdr)
    run = p.add_run(text)
    set_run_font(run, size=16 if level == 1 else 13, bold=True, color=RED if level == 1 else RED2)
    return p


def add_bullet(doc, text, size=11):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Cm(0.75)
    r = p.add_run(text)
    set_run_font(r, size=size, color=MUT)
    return p


def callout(doc, label, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_width(table, 16.5)
    cell = table.cell(0, 0)
    shade_cell(cell, SOFT)
    set_cell_borders(cell, RED, "12")
    set_cell_margins(cell, 100, 100, 140, 140)
    cell.text = ""
    p1 = cell.paragraphs[0]
    p1.paragraph_format.space_after = Pt(4)
    r1 = p1.add_run(label.upper())
    set_run_font(r1, size=9, bold=True, color=RED)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run(text)
    set_run_font(r2, size=11, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)


def make_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_width(table, 16.5)
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        shade_cell(cell, RED)
        set_cell_borders(cell, RED, "4")
        set_cell_margins(cell, 70, 70, 90, 90)
        cell_para(cell, h, size=9, bold=True, color=WHITE, space_after=0)
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            cell = table.rows[r_i + 1].cells[c_i]
            shade_cell(cell, CREAM if r_i % 2 else WHITE)
            set_cell_borders(cell, LINE, "4")
            set_cell_margins(cell, 70, 70, 90, 90)
            cell_para(cell, val, size=10, bold=c_i == 0, space_after=0)
        prevent_row_split(table.rows[r_i + 1])
    prevent_row_split(table.rows[0])
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Cm(w)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(8)
    return table


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def html_table(headers, rows):
    head = "".join(f"<th>{escape(h)}</th>" for h in headers)
    body = []
    for row in rows:
        cells = "".join(
            f'<td style="{"font-weight:700;" if i == 0 else ""}">{escape(val)}</td>'
            for i, val in enumerate(row)
        )
        body.append(f"<tr>{cells}</tr>")
    return (
        '<table class="grid" border="1" cellspacing="0" cellpadding="0">'
        f"<thead><tr>{head}</tr></thead><tbody>{''.join(body)}</tbody></table>"
    )


def html_callout(label, text):
    return (
        '<table class="callout" border="0" cellspacing="0" cellpadding="0"><tr><td>'
        f'<p class="callout-label">{escape(label.upper())}</p>'
        f"<p>{escape(text)}</p>"
        "</td></tr></table>"
    )


def build_docx():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(2.0)

    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = hp.add_run("CID — Clínica de Serviços Médicos  ×  V4 Company  ·  Proposta comercial")
    set_run_font(r, size=8, color=RED, bold=True)

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = fp.add_run("Empresa que Vende + SEO Growth  ·  Confidencial  ·  Página ")
    set_run_font(r, size=8, color=MUT)
    add_page_number(fp)

    cover = doc.add_table(rows=1, cols=1)
    set_table_width(cover, 17.0)
    c = cover.cell(0, 0)
    shade_cell(c, DARK)
    set_cell_borders(c, DARK, "0")
    set_cell_margins(c, 280, 280, 220, 220)
    c.text = ""
    p = c.paragraphs[0]
    r = p.add_run("PROPOSTA COMERCIAL")
    set_run_font(r, size=11, bold=True, color="F5B5B8")
    p2 = c.add_paragraph()
    p2.paragraph_format.space_before = Pt(10)
    r = p2.add_run("Empresa que Vende + SEO Growth")
    set_run_font(r, size=28, bold=True, color=WHITE)
    p3 = c.add_paragraph()
    p3.paragraph_format.space_after = Pt(12)
    r = p3.add_run("CID — Clínica de Serviços Médicos  ×  V4 Company")
    set_run_font(r, size=16, color="FFD0D0")
    p4 = c.add_paragraph()
    r = p4.add_run(
        "Consultoria, acompanhamento e execução assistida + origem orgânica. "
        "Ciclo de 6 meses no cartão de crédito. Sem verba de mídia paga."
    )
    set_run_font(r, size=12, color="E8D0D0")

    add_p(doc, "", size=8, space_after=4)
    make_table(
        doc,
        ["Campo", "Informação"],
        [
            ["Cliente", "CID — Clínica de Serviços Médicos"],
            ["Produto 1", "Empresa que Vende (EQV) — execução assistida"],
            ["Produto 2", "SEO Growth — conteúdo + links + previsibilidade (6 meses)"],
            ["Prazo", "6 meses  ·  5 fases  ·  squad (Conteúdo, SEO Growth, Comercial, CRM)"],
            ["Separados", "EQV R$ 35.000,00  ·  SEO Growth R$ 21.600,00 (6× R$ 3.600)"],
            ["Ambos juntos (soma)", "6× R$ 9.433,33 no cartão  ·  TCV R$ 56.600,00"],
            ["Pacote · condição comercial", "6× R$ 5.418,57 no cartão de crédito  ·  TCV R$ 32.511,42"],
            ["Condução", "V4 Company  ·  consultor estratégico (1h/mês) + squad de execução assistida"],
            ["Incluso", "Comunidade, monitoria coletiva mensal, conteúdo semanal e Fábrica de Receita (imersão presencial anual)"],
            ["Garantia", "30 dias de serviço — reembolso incondicional se a entrega não fizer sentido"],
        ],
        [4.6, 11.9],
    )

    callout(
        doc,
        "Leitura central",
        "A V4 não executa no lugar da CID: estrutura o negócio, treina o time e acompanha a execução. "
        "Quem opera é a CID. A máquina fica na empresa quando o ciclo acaba. "
        "A ordem é clareza → sistema → execução. SEO Growth entra como origem complementar, sem mídia paga.",
    )

    add_heading_styled(doc, "1. Objetivo")
    add_p(
        doc,
        "Instalar, em 6 meses, o método para a CID vender com previsibilidade: diagnóstico contínuo "
        "de maturidade GTM, consultoria estratégica mensal, execução assistida (conteúdo, comercial e CRM) "
        "e origem complementar em SEO Growth — sem verba de mídia paga. Comunidade de pares e imersão "
        "presencial da Fábrica de Receita.",
        align="justify",
    )
    add_p(
        doc,
        "O resultado esperado não é um PDF para a gaveta nem uma agência que some quando o contrato acaba. "
        "É o time operando com tese, artefato e cadência: conteúdo que gera demanda, comercial que fecha, "
        "CRM que governa o pipeline, e a CID aparecendo quando RH e Compras pesquisam — inclusive quando "
        "perguntam para a IA.",
        align="justify",
    )

    add_heading_styled(doc, "2. O problema que o programa resolve")
    make_table(
        doc,
        ["Sintoma", "Efeito na operação"],
        [
            [
                "Receita imprevisível / concentração",
                "O mês é esforço, não método. 90% em um contrato não tem plano B privado.",
            ],
            [
                "Processo na cabeça de uma pessoa",
                "O comercial é o sócio Leonardo. Sem sistema, não se repete.",
            ],
            [
                "Marketing no achismo / nunca investiu em mídia",
                "Canal, conteúdo e origem sem tese. Mídia paga não cabe na margem.",
            ],
            [
                "Agência executa — você não aprende",
                "Quando o contrato acaba, o resultado vai junto.",
            ],
            [
                "Comercial amador",
                "Script, follow-up e pipeline improvisados. A proposta não tem business case.",
            ],
            [
                "Posicionamento frouxo",
                "O mercado vê mão de obra avulsa, não gestão e continuidade.",
            ],
        ],
        [5.5, 11.0],
    )

    add_heading_styled(doc, "3. A solução")
    add_p(
        doc,
        "A ordem de construção é clareza → sistema → execução. Sem método, origem vira aposta. "
        "Com método, receita vira consequência.",
        align="justify",
    )
    make_table(
        doc,
        ["Pilar", "O que instala"],
        [
            ["01 · Diagnóstico contínuo", "Maturidade GTM monitorada. Gargalos identificados a cada ciclo."],
            ["02 · Consultoria estratégica", "1h/mês com consultor V4: prioridades e créditos do squad."],
            ["03 · Execução assistida", "Squad: Conteúdo, SEO Growth, Comercial e CRM. Quem opera é a CID."],
            ["04 · Comunidade", "Peers, lives mensais e grupo fechado no WhatsApp."],
            ["05 · Fábrica de Receita", "Imersão presencial anual: hot seats, cases e workshops."],
        ],
        [5.2, 11.3],
    )

    add_heading_styled(doc, "4. Jornada de 6 meses")
    make_table(
        doc,
        ["Fase", "Quando", "Nome", "Função"],
        [
            ["01", "Mês 1", "Onboarding", "Kickoff, maturidade GTM e prioridades do ciclo."],
            ["02", "Mês 2", "EE 3.0 completa", "3P-IA, posicionamento ativo e canal de vendas no ar."],
            ["03", "Mês 3", "Consultoria de conteúdo", "Produção assistida + mensagem para RH / Compras / Financeiro."],
            ["04", "Mês 4", "SEO Growth", "Base técnica, intenção, 4 posts/mês e entidade para busca e IA."],
            ["05", "Mês 5/6", "Consultoria comercial + CRM", "Scripts, follow-up, pipeline e monetização do funil."],
        ],
        [2.0, 2.6, 4.6, 7.3],
    )

    add_heading_styled(doc, "5. Três sessões com entregáveis")
    add_heading_styled(doc, "Sessão 01 — Produto, oferta, posicionamento, conteúdo", 2)
    add_p(
        doc,
        "O especialista V4 orienta a estratégia, calibra a mensagem para o ICP (RH, Compras, Financeiro) "
        "e acompanha a execução mês a mês. A CID para de falar de si e passa a falar da dor do comprador.",
        align="justify",
    )
    for item in [
        "Oferta e posicionamento de gestão — não de mão de obra",
        "Proposta com business case, cronograma e governança",
        "Linha editorial e mensagem por persona",
        "Prova comercial: SLA, indicador, caso autorizado",
    ]:
        add_bullet(doc, item)

    add_heading_styled(doc, "Sessão 02 — SEO Growth", 2)
    add_p(
        doc,
        "SEO Growth é o plano de origem orgânica: crescimento consistente com conteúdo, links e "
        "previsibilidade. Não é post solto, não é “mexer no site”, não é garantia de posição. "
        "GEO não é um módulo à parte — é o efeito de conteúdo factual e entidade: RH e Compras já "
        "perguntam para a IA quem contratar; a CID precisa ser citada como referência de gestão.",
        align="justify",
    )
    add_p(doc, "Incluso no ciclo de 6 meses:", bold=True, size=11, space_after=4)
    for item in [
        "Base técnica, Search Console, Google Business Profile e otimização inicial",
        "Auditoria técnica contínua",
        "Pesquisa avançada sem limite + intenção de busca do comprador",
        "4 blogposts/mês + briefings + snippet",
        "Link building leve (1–2/mês) e recuperação de links quebrados",
        "SEO local básico, se aplicável à operação",
        "Gestão quinzenal + call mensal",
    ]:
        add_bullet(doc, item)
    add_p(doc, "Fora deste recorte:", bold=True, size=11, space_before=6, space_after=4)
    for item in [
        "Cluster pilar + satélites",
        "Link building estratégico / Digital PR",
        "Core Web Vitals avançado",
        "CRO orgânico",
    ]:
        add_bullet(doc, item)

    add_heading_styled(doc, "Sessão 03 — Comercial e CRM", 2)
    add_p(
        doc,
        "Acompanhamento comercial próximo: o especialista revisa scripts, analisa as objeções reais "
        "do mês e ajusta o processo com o time. O comercial deixa de morar na cabeça do Leonardo.",
        align="justify",
    )
    for item in [
        "Revisão dos scripts de venda e cliente oculto",
        "Análise de objeções reais do mês e treinamento comercial",
        "Réguas de conversão, follow-up, renovação e upsell na base",
        "CRM ativo: cadência, pipeline e próxima ação — não agenda",
    ]:
        add_bullet(doc, item)

    add_heading_styled(doc, "6. Recorte EQV × SEO Growth")
    make_table(
        doc,
        ["Frente", "O que entrega"],
        [
            ["EQV · sessão 01", "Mensagem, oferta, prova e linha editorial. O que dizer e para quem."],
            ["SEO Growth", "Produção orgânica contínua: 4 posts/mês, técnica, intenção e links leves."],
            ["EQV · sessão 03", "Scripts, follow-up, pipeline e CRM. Como fechar e governar."],
        ],
        [4.8, 11.7],
    )
    callout(
        doc,
        "Sem sobreposição de entrega",
        "O EQV define a tese comercial. O SEO Growth produz o ativo orgânico com essa tese. "
        "Um não substitui o outro: sem mensagem, o post não converte; sem produção, a mensagem não aparece na busca.",
    )

    add_heading_styled(doc, "7. Investimento")
    price = doc.add_table(rows=1, cols=2)
    set_table_width(price, 16.5)
    left, right = price.cell(0, 0), price.cell(0, 1)
    shade_cell(left, DARK)
    shade_cell(right, SOFT)
    set_cell_borders(left, DARK, "4")
    set_cell_borders(right, RED, "8")
    set_cell_margins(left, 160, 160, 160, 160)
    set_cell_margins(right, 140, 140, 140, 140)
    left.text = ""
    p = left.paragraphs[0]
    r = p.add_run("VALOR DE REFERÊNCIA")
    set_run_font(r, size=8, bold=True, color="F5B5B8")
    p2 = left.add_paragraph()
    r = p2.add_run("6× R$ 9.433,33")
    set_run_font(r, size=16, color="C9A0A0")
    p2.runs[0].font.strike = True
    p3 = left.add_paragraph()
    r = p3.add_run("CONDIÇÃO COMERCIAL CID")
    set_run_font(r, size=8, bold=True, color="F5B5B8")
    p4 = left.add_paragraph()
    r = p4.add_run("6× R$ 5.418,57")
    set_run_font(r, size=26, bold=True, color=WHITE)
    p5 = left.add_paragraph()
    r = p5.add_run("EQV execução assistida + SEO Growth\n6 meses no cartão de crédito")
    set_run_font(r, size=10, color="E8D0D0")
    right.text = ""
    p = right.paragraphs[0]
    r = p.add_run("CONDIÇÕES")
    set_run_font(r, size=8, bold=True, color=RED)
    add_cell_run(right, "6× R$ 5.418,57 no cartão de crédito", size=14, bold=True, color=INK)
    add_cell_run(right, "TCV do pacote: R$ 32.511,42.", size=10, color=MUT)
    add_cell_run(
        right,
        "Se fossem separados: EQV R$ 35.000,00 + SEO Growth R$ 21.600,00 (6× R$ 3.600) = R$ 56.600,00.",
        size=10,
        color=MUT,
    )
    add_cell_run(
        right,
        "Ciclo de 6 meses. A máquina fica no time. Não é aluguel de agência. Mídia paga não entra neste desenho.",
        size=10,
        color=MUT,
    )
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(10)

    add_heading_styled(doc, "7.1 Se fossem contratados separados", 2)
    make_table(
        doc,
        ["Frente", "Valor"],
        [
            ["Execução assistida (EQV)", "R$ 35.000,00"],
            ["SEO Growth (6× R$ 3.600)", "R$ 21.600,00"],
            ["Ambos juntos (soma)", "6× R$ 9.433,33 no cartão  ·  TCV R$ 56.600,00"],
            ["Pacote · condição comercial", "6× R$ 5.418,57 no cartão  ·  TCV R$ 32.511,42"],
        ],
        [8.0, 8.5],
    )

    add_heading_styled(doc, "8. Garantia")
    add_p(
        doc,
        "Se após 30 dias de serviço você não estiver satisfeito com a entrega, devolvemos o "
        "investimento efetuado. Risco zero para entrar. O método precisa fazer sentido na operação — "
        "não só no papel.",
        align="justify",
    )

    add_heading_styled(doc, "9. Próximos passos")
    make_table(
        doc,
        ["Responsável", "Ação"],
        [
            [
                "CID (Leonardo)",
                "Validar a condição comercial (6× R$ 5.418,57 no cartão) e formalizar o start.",
            ],
            [
                "V4",
                "Kickoff do mês 01: diagnóstico de maturidade GTM, prioridades do ciclo e alocação do squad.",
            ],
            [
                "V4 + CID",
                "Entrar na comunidade, na monitoria mensal e no calendário da Fábrica de Receita a partir da assinatura.",
            ],
        ],
        [4.4, 12.1],
    )
    add_p(
        doc,
        "Documento elaborado pela V4 Company para a CID — Clínica de Serviços Médicos. "
        "Valores e prazos conforme condição comercial do pacote. Confidencial.",
        size=9,
        color=MUT,
        italic=True,
        space_before=12,
    )
    add_p(
        doc,
        "CID — Clínica de Serviços Médicos × V4 Company  ·  Empresa que Vende + SEO Growth  ·  Confidencial",
        size=9,
        color=RED,
        bold=True,
        align="center",
        space_before=8,
    )
    OUT_DIR.mkdir(exist_ok=True)
    doc.save(OUT_DOCX)
    print("Wrote", OUT_DOCX)


def build_html():
    def h1(t):
        return f"<h1>{escape(t)}</h1>"

    def h2(t):
        return f"<h2>{escape(t)}</h2>"

    def p(t, cls="justify"):
        return f'<p class="{cls}">{escape(t)}</p>'

    def ul(items):
        return "<ul>" + "".join(f"<li>{escape(i)}</li>" for i in items) + "</ul>"

    parts = [
        """<table class="cover" border="0" cellspacing="0" cellpadding="0"><tr><td>
<p class="cover-kicker">PROPOSTA COMERCIAL</p>
<p class="cover-title">Empresa que Vende + SEO Growth</p>
<p class="cover-sub">CID — Clínica de Serviços Médicos  ×  V4 Company</p>
<p class="cover-lead">Consultoria, acompanhamento e execução assistida + origem orgânica. Ciclo de 6 meses no cartão de crédito. Sem verba de mídia paga.</p>
</td></tr></table>""",
        html_table(
            ["Campo", "Informação"],
            [
                ["Cliente", "CID — Clínica de Serviços Médicos"],
                ["Produto 1", "Empresa que Vende (EQV) — execução assistida"],
                ["Produto 2", "SEO Growth — conteúdo + links + previsibilidade (6 meses)"],
                ["Prazo", "6 meses  ·  5 fases  ·  squad (Conteúdo, SEO Growth, Comercial, CRM)"],
                ["Separados", "EQV R$ 35.000,00  ·  SEO Growth R$ 21.600,00 (6× R$ 3.600)"],
                ["Ambos juntos (soma)", "6× R$ 9.433,33 no cartão  ·  TCV R$ 56.600,00"],
                ["Pacote · condição comercial", "6× R$ 5.418,57 no cartão de crédito  ·  TCV R$ 32.511,42"],
                ["Condução", "V4 Company  ·  consultor estratégico (1h/mês) + squad de execução assistida"],
                ["Incluso", "Comunidade, monitoria coletiva mensal, conteúdo semanal e Fábrica de Receita (imersão presencial anual)"],
                ["Garantia", "30 dias de serviço — reembolso incondicional se a entrega não fizer sentido"],
            ],
        ),
        html_callout(
            "Leitura central",
            "A V4 não executa no lugar da CID: estrutura o negócio, treina o time e acompanha a execução. "
            "Quem opera é a CID. A máquina fica na empresa quando o ciclo acaba. "
            "A ordem é clareza → sistema → execução. SEO Growth entra como origem complementar, sem mídia paga.",
        ),
        h1("1. Objetivo"),
        p(
            "Instalar, em 6 meses, o método para a CID vender com previsibilidade: diagnóstico contínuo "
            "de maturidade GTM, consultoria estratégica mensal, execução assistida (conteúdo, comercial e CRM) "
            "e origem complementar em SEO Growth — sem verba de mídia paga. Comunidade de pares e imersão "
            "presencial da Fábrica de Receita."
        ),
        p(
            "O resultado esperado não é um PDF para a gaveta nem uma agência que some quando o contrato acaba. "
            "É o time operando com tese, artefato e cadência: conteúdo que gera demanda, comercial que fecha, "
            "CRM que governa o pipeline, e a CID aparecendo quando RH e Compras pesquisam — inclusive quando "
            "perguntam para a IA."
        ),
        h1("2. O problema que o programa resolve"),
        html_table(
            ["Sintoma", "Efeito na operação"],
            [
                ["Receita imprevisível / concentração", "O mês é esforço, não método. 90% em um contrato não tem plano B privado."],
                ["Processo na cabeça de uma pessoa", "O comercial é o sócio Leonardo. Sem sistema, não se repete."],
                ["Marketing no achismo / nunca investiu em mídia", "Canal, conteúdo e origem sem tese. Mídia paga não cabe na margem."],
                ["Agência executa — você não aprende", "Quando o contrato acaba, o resultado vai junto."],
                ["Comercial amador", "Script, follow-up e pipeline improvisados. A proposta não tem business case."],
                ["Posicionamento frouxo", "O mercado vê mão de obra avulsa, não gestão e continuidade."],
            ],
        ),
        h1("3. A solução"),
        p("A ordem de construção é clareza → sistema → execução. Sem método, origem vira aposta. Com método, receita vira consequência."),
        html_table(
            ["Pilar", "O que instala"],
            [
                ["01 · Diagnóstico contínuo", "Maturidade GTM monitorada. Gargalos identificados a cada ciclo."],
                ["02 · Consultoria estratégica", "1h/mês com consultor V4: prioridades e créditos do squad."],
                ["03 · Execução assistida", "Squad: Conteúdo, SEO Growth, Comercial e CRM. Quem opera é a CID."],
                ["04 · Comunidade", "Peers, lives mensais e grupo fechado no WhatsApp."],
                ["05 · Fábrica de Receita", "Imersão presencial anual: hot seats, cases e workshops."],
            ],
        ),
        h1("4. Jornada de 6 meses"),
        html_table(
            ["Fase", "Quando", "Nome", "Função"],
            [
                ["01", "Mês 1", "Onboarding", "Kickoff, maturidade GTM e prioridades do ciclo."],
                ["02", "Mês 2", "EE 3.0 completa", "3P-IA, posicionamento ativo e canal de vendas no ar."],
                ["03", "Mês 3", "Consultoria de conteúdo", "Produção assistida + mensagem para RH / Compras / Financeiro."],
                ["04", "Mês 4", "SEO Growth", "Base técnica, intenção, 4 posts/mês e entidade para busca e IA."],
                ["05", "Mês 5/6", "Consultoria comercial + CRM", "Scripts, follow-up, pipeline e monetização do funil."],
            ],
        ),
        h1("5. Três sessões com entregáveis"),
        h2("Sessão 01 — Produto, oferta, posicionamento, conteúdo"),
        p(
            "O especialista V4 orienta a estratégia, calibra a mensagem para o ICP (RH, Compras, Financeiro) "
            "e acompanha a execução mês a mês. A CID para de falar de si e passa a falar da dor do comprador."
        ),
        ul(
            [
                "Oferta e posicionamento de gestão — não de mão de obra",
                "Proposta com business case, cronograma e governança",
                "Linha editorial e mensagem por persona",
                "Prova comercial: SLA, indicador, caso autorizado",
            ]
        ),
        h2("Sessão 02 — SEO Growth"),
        p(
            "SEO Growth é o plano de origem orgânica: crescimento consistente com conteúdo, links e "
            "previsibilidade. Não é post solto, não é “mexer no site”, não é garantia de posição. "
            "GEO não é um módulo à parte — é o efeito de conteúdo factual e entidade: RH e Compras já "
            "perguntam para a IA quem contratar; a CID precisa ser citada como referência de gestão."
        ),
        '<p class="label">Incluso no ciclo de 6 meses</p>',
        ul(
            [
                "Base técnica, Search Console, Google Business Profile e otimização inicial",
                "Auditoria técnica contínua",
                "Pesquisa avançada sem limite + intenção de busca do comprador",
                "4 blogposts/mês + briefings + snippet",
                "Link building leve (1–2/mês) e recuperação de links quebrados",
                "SEO local básico, se aplicável à operação",
                "Gestão quinzenal + call mensal",
            ]
        ),
        '<p class="label">Fora deste recorte</p>',
        ul(
            [
                "Cluster pilar + satélites",
                "Link building estratégico / Digital PR",
                "Core Web Vitals avançado",
                "CRO orgânico",
            ]
        ),
        h2("Sessão 03 — Comercial e CRM"),
        p(
            "Acompanhamento comercial próximo: o especialista revisa scripts, analisa as objeções reais "
            "do mês e ajusta o processo com o time. O comercial deixa de morar na cabeça do Leonardo."
        ),
        ul(
            [
                "Revisão dos scripts de venda e cliente oculto",
                "Análise de objeções reais do mês e treinamento comercial",
                "Réguas de conversão, follow-up, renovação e upsell na base",
                "CRM ativo: cadência, pipeline e próxima ação — não agenda",
            ]
        ),
        h1("6. Recorte EQV × SEO Growth"),
        html_table(
            ["Frente", "O que entrega"],
            [
                ["EQV · sessão 01", "Mensagem, oferta, prova e linha editorial. O que dizer e para quem."],
                ["SEO Growth", "Produção orgânica contínua: 4 posts/mês, técnica, intenção e links leves."],
                ["EQV · sessão 03", "Scripts, follow-up, pipeline e CRM. Como fechar e governar."],
            ],
        ),
        html_callout(
            "Sem sobreposição de entrega",
            "O EQV define a tese comercial. O SEO Growth produz o ativo orgânico com essa tese. "
            "Um não substitui o outro: sem mensagem, o post não converte; sem produção, a mensagem não aparece na busca.",
        ),
        h1("7. Investimento"),
        """<table class="price" border="0" cellspacing="0" cellpadding="0">
<tr>
<td class="price-left" width="50%">
<p class="cover-kicker">VALOR DE REFERÊNCIA</p>
<p class="strike">6× R$ 9.433,33</p>
<p class="cover-kicker">CONDIÇÃO COMERCIAL CID</p>
<p class="price-now">6× R$ 5.418,57</p>
<p class="cover-lead">EQV execução assistida + SEO Growth<br>6 meses no cartão de crédito</p>
</td>
<td class="price-right" width="50%">
<p class="callout-label">CONDIÇÕES</p>
<p class="price-month">6× R$ 5.418,57 no cartão de crédito</p>
<p>TCV do pacote: R$ 32.511,42.</p>
<p>Se fossem separados: EQV R$ 35.000,00 + SEO Growth R$ 21.600,00 (6× R$ 3.600) = R$ 56.600,00.</p>
<p>Ciclo de 6 meses. A máquina fica no time. Não é aluguel de agência. Mídia paga não entra neste desenho.</p>
</td>
</tr>
</table>""",
        h2("7.1 Se fossem contratados separados"),
        html_table(
            ["Frente", "Valor"],
            [
                ["Execução assistida (EQV)", "R$ 35.000,00"],
                ["SEO Growth (6× R$ 3.600)", "R$ 21.600,00"],
                ["Ambos juntos (soma)", "6× R$ 9.433,33 no cartão  ·  TCV R$ 56.600,00"],
                ["Pacote · condição comercial", "6× R$ 5.418,57 no cartão  ·  TCV R$ 32.511,42"],
            ],
        ),
        h1("8. Garantia"),
        p(
            "Se após 30 dias de serviço você não estiver satisfeito com a entrega, devolvemos o "
            "investimento efetuado. Risco zero para entrar. O método precisa fazer sentido na operação — "
            "não só no papel."
        ),
        h1("9. Próximos passos"),
        html_table(
            ["Responsável", "Ação"],
            [
                ["CID (Leonardo)", "Validar a condição comercial (6× R$ 5.418,57 no cartão) e formalizar o start."],
                ["V4", "Kickoff do mês 01: diagnóstico de maturidade GTM, prioridades do ciclo e alocação do squad."],
                ["V4 + CID", "Entrar na comunidade, na monitoria mensal e no calendário da Fábrica de Receita a partir da assinatura."],
            ],
        ),
        p(
            "Documento elaborado pela V4 Company para a CID — Clínica de Serviços Médicos. "
            "Valores e prazos conforme condição comercial do pacote. Confidencial.",
            "muted",
        ),
        '<p class="footer-brand">CID — Clínica de Serviços Médicos × V4 Company  ·  Empresa que Vende + SEO Growth  ·  Confidencial</p>',
    ]

    body = "\n".join(parts)
    html = f"""<!DOCTYPE html>
<html xmlns:o="urn:schemas-microsoft-com:office:office"
      xmlns:w="urn:schemas-microsoft-com:office:word"
      xmlns="http://www.w3.org/TR/REC-html40">
<head>
<meta charset="utf-8">
<meta http-equiv="Content-Type" content="text/html; charset=utf-8">
<meta name="ProgId" content="Word.Document">
<meta name="Generator" content="Microsoft Word 15">
<title>CID × V4 — Empresa que Vende + SEO Growth</title>
<!--[if gte mso 9]>
<xml>
  <w:WordDocument>
    <w:View>Print</w:View>
    <w:Zoom>100</w:Zoom>
    <w:DoNotOptimizeForBrowser/>
  </w:WordDocument>
</xml>
<![endif]-->
<style>
  @page WordSection1 {{ size: 21cm 29.7cm; margin: 1.8cm 2cm 2cm 2cm; }}
  div.WordSection1 {{ page: WordSection1; }}
  body {{
    font-family: Calibri, Arial, sans-serif; font-size: 11pt; color: #1A1A1A;
    line-height: 1.35; margin: 24px auto; max-width: 820px; background: #f4f0ef;
  }}
  div.WordSection1 {{ background: #fff; padding: 28px 32px 48px; }}
  h1 {{ font-size: 16pt; color: #A10F14; border-bottom: 1.5pt solid #A10F14; padding-bottom: 4pt; margin: 18pt 0 8pt; }}
  h2 {{ font-size: 12.5pt; color: #E50914; margin: 12pt 0 6pt; }}
  p {{ margin: 0 0 8pt; }}
  p.justify {{ text-align: justify; }}
  p.muted {{ color: #4A4A4A; font-size: 10.5pt; font-style: italic; }}
  p.label {{ color: #E50914; font-weight: 700; margin: 8pt 0 4pt; }}
  p.footer-brand {{ color: #A10F14; font-weight: 700; text-align: center; font-size: 9pt; margin-top: 16pt; }}
  ul {{ margin: 4pt 0 10pt 18pt; padding: 0; }}
  li {{ margin: 0 0 3pt; color: #4A4A4A; }}
  table.grid {{ width: 100%; border-collapse: collapse; margin: 8pt 0 12pt; font-size: 10pt; }}
  table.grid th {{ background: #A10F14; color: #fff; text-align: left; font-size: 9pt; padding: 6pt 8pt; border: 1pt solid #A10F14; }}
  table.grid td {{ border: 1pt solid #E6D4D4; padding: 6pt 8pt; vertical-align: top; }}
  table.callout {{ width: 100%; margin: 8pt 0 12pt; border-left: 4pt solid #A10F14; background: #F8EFEF; }}
  table.callout td {{ padding: 10pt 12pt; }}
  .callout-label {{ color: #A10F14; font-size: 8.5pt; font-weight: 700; letter-spacing: .08em; margin: 0 0 4pt; }}
  table.cover {{ width: 100%; background: #280001; margin: 0 0 14pt; }}
  table.cover td {{ padding: 22pt 20pt; }}
  .cover-kicker {{ color: #F5B5B8; font-size: 9pt; font-weight: 700; letter-spacing: .12em; margin: 0 0 8pt; }}
  .cover-title {{ color: #fff; font-size: 26pt; font-weight: 700; margin: 0 0 6pt; }}
  .cover-sub {{ color: #FFD0D0; font-size: 14pt; margin: 0 0 10pt; }}
  .cover-lead {{ color: #E8D0D0; font-size: 11pt; margin: 0; }}
  table.price {{ width: 100%; margin: 8pt 0 12pt; }}
  td.price-left {{ background: #280001; color: #fff; padding: 16pt; vertical-align: top; width: 50%; }}
  td.price-right {{ background: #F8EFEF; border: 1.5pt solid #A10F14; padding: 14pt; vertical-align: top; width: 50%; }}
  .strike {{ color: #C9A0A0; text-decoration: line-through; font-size: 14pt; margin: 0 0 8pt; }}
  .price-now {{ color: #fff; font-size: 24pt; font-weight: 700; margin: 0 0 8pt; }}
  .price-month {{ font-size: 13pt; font-weight: 700; margin: 0 0 8pt; }}
</style>
</head>
<body>
<div class="WordSection1">
{body}
</div>
</body>
</html>
"""
    OUT_DIR.mkdir(exist_ok=True)
    OUT_HTML.write_text(html, encoding="utf-8")
    OUT_DOC.write_text(html, encoding="utf-8")
    print("Wrote", OUT_HTML)
    print("Wrote", OUT_DOC)


if __name__ == "__main__":
    build_docx()
    build_html()
